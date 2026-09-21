from openpyxl import load_workbook
from pymongo import UpdateOne

from auth import hash_password
from db import courses_collection, students_collection

REGISTRATION_SHEET = "Overall Regitrations"
COURSE_SHEET = "courses"

SEMESTERS = {
    "2-1": {
        "course_code": 8,
        "course_name": 9,
        "erp_status": 10,
        "current_status": 11,
        "verifiable_link": 12,
        "replacement_code": 13,
        "replacement_name": 14
    },
    "2-2": {
        "course_code": 15,
        "course_name": 16,
        "erp_status": 17,
        "current_status": 18,
        "verifiable_link": 19,
        "replacement_code": 20,
        "replacement_name": 21
    },
    "3-1": {
        "course_code": 22,
        "course_name": 23,
        "erp_status": 24,
        "current_status": 25,
        "verifiable_link": 26,
        "replacement_code": 27,
        "replacement_name": 28
    },
    "3-2": {
        "course_code": 29,
        "course_name": 30,
        "erp_status": 31,
        "current_status": 32,
        "verifiable_link": 33,
        "replacement_code": 34,
        "replacement_name": 35
    }
}


def clean(value):
    if value is None:
        return ""

    return str(value).replace("\xa0", " ").strip()


def import_workbook(path):
    """
    Parse the master Excel workbook and upsert its data into MongoDB.

    Existing students keep their password/login state; only their
    academic data is refreshed. New students are provisioned with a
    default password equal to their student ID.
    """

    wb = load_workbook(path, data_only=False)

    if REGISTRATION_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{REGISTRATION_SHEET}' not found in workbook."
        )

    if COURSE_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{COURSE_SHEET}' not found in workbook."
        )

    # --------------------------------------------------------
    # Courses
    # --------------------------------------------------------

    course_ws = wb[COURSE_SHEET]

    course_ops = []
    course_count = 0

    for row in range(2, course_ws.max_row + 1):

        code = clean(course_ws.cell(row, 1).value)
        title = clean(course_ws.cell(row, 2).value)

        if code and title:
            course_ops.append(
                UpdateOne(
                    {"code": code},
                    {"$set": {"code": code, "title": title}},
                    upsert=True
                )
            )
            course_count += 1

    if course_ops:
        courses_collection.bulk_write(course_ops)

    # --------------------------------------------------------
    # Students
    # --------------------------------------------------------

    reg_ws = wb[REGISTRATION_SHEET]

    existing_students = {
        doc["student_id"]: doc
        for doc in students_collection.find(
            {}, {"student_id": 1, "password_hash": 1,
                 "must_change_password": 1, "history": 1}
        )
    }

    student_ops = []
    student_count = 0

    for row in range(2, reg_ws.max_row + 1):

        student_id = clean(reg_ws.cell(row, 3).value)

        if not student_id:
            continue

        existing = existing_students.get(student_id)
        existing_history = (existing or {}).get("history", {})

        history = {}

        for semester, columns in SEMESTERS.items():

            replacement_code = clean(
                reg_ws.cell(row, columns["replacement_code"]).value
            )

            existing_semester = existing_history.get(semester, {})

            if "replacement_locked" in existing_semester:
                # Preserve admin lock/unlock decisions across re-imports.
                replacement_locked = existing_semester["replacement_locked"]
            else:
                replacement_locked = bool(replacement_code)

            history[semester] = {
                "course_code":
                    clean(reg_ws.cell(row, columns["course_code"]).value),
                "course_name":
                    clean(reg_ws.cell(row, columns["course_name"]).value),
                "erp_status":
                    clean(reg_ws.cell(row, columns["erp_status"]).value),
                "current_status":
                    clean(reg_ws.cell(row, columns["current_status"]).value),
                "verifiable_link":
                    clean(reg_ws.cell(row, columns["verifiable_link"]).value),
                "replacement_code": replacement_code,
                "replacement_name":
                    clean(
                        reg_ws.cell(row, columns["replacement_name"]).value
                    ),
                "replacement_locked": replacement_locked
            }

        document = {
            "student_id": student_id,
            "student_name": clean(reg_ws.cell(row, 4).value),
            "faculty_incharge": clean(reg_ws.cell(row, 2).value),
            "counsellor_emp_id": clean(reg_ws.cell(row, 5).value),
            "counsellor_name": clean(reg_ws.cell(row, 6).value),
            "placed": clean(reg_ws.cell(row, 7).value),
            "required": reg_ws.cell(row, 36).value,
            "total_registered": reg_ws.cell(row, 37).value,
            "gender": clean(reg_ws.cell(row, 38).value),
            "current_4_1": clean(reg_ws.cell(row, 39).value),
            "registered_4_1": bool(reg_ws.cell(row, 40).value),
            "current_4_2": clean(reg_ws.cell(row, 41).value),
            "history": history
        }

        if existing and existing.get("password_hash"):
            document["password_hash"] = existing["password_hash"]
            document["must_change_password"] = existing.get(
                "must_change_password", True
            )
        else:
            document["password_hash"] = hash_password(student_id)
            document["must_change_password"] = True

        student_ops.append(
            UpdateOne(
                {"student_id": student_id},
                {"$set": document},
                upsert=True
            )
        )
        student_count += 1

    if student_ops:
        for i in range(0, len(student_ops), 500):
            students_collection.bulk_write(student_ops[i:i + 500])

    return {
        "courses_imported": course_count,
        "students_imported": student_count
    }
