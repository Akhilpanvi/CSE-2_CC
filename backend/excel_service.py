
from pathlib import Path
from openpyxl import load_workbook


# ============================================================
# FILE CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = (
    BASE_DIR
    / "data"
    / "Y23 Skill Certification Details_ 10-08-26.xlsx"
)

REGISTRATION_SHEET = "Overall Regitrations"
COURSE_SHEET = "courses"


# ============================================================
# COMMON HELPERS
# ============================================================

def clean(value):
    """
    Convert Excel values into clean strings.
    """
    if value is None:
        return ""

    return str(value).replace("\xa0", " ").strip()


def get_workbook():
    """
    Load the Excel workbook.
    """
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    return load_workbook(
        EXCEL_FILE,
        data_only=False
    )


# ============================================================
# COURSES
# ============================================================

def get_courses():

    wb = get_workbook()

    if COURSE_SHEET not in wb.sheetnames:
        wb.close()
        return []

    ws = wb[COURSE_SHEET]

    courses = []

    for row in range(2, ws.max_row + 1):

        code = clean(
            ws.cell(row, 1).value
        )

        title = clean(
            ws.cell(row, 2).value
        )

        if code and title:

            courses.append({
                "code": code,
                "title": title
            })

    wb.close()

    return courses


# ============================================================
# FIND STUDENT
# ============================================================

def find_student(student_id):

    wb = get_workbook()

    ws = wb[REGISTRATION_SHEET]

    student_id = clean(student_id)

    for row in range(2, ws.max_row + 1):

        # STUDENT_ID is column 3
        current_id = clean(
            ws.cell(row, 3).value
        )

        if current_id == student_id:

            student = {

                "row": row,

                "student_id":
                    current_id,

                "student_name":
                    clean(
                        ws.cell(row, 4).value
                    ),

                # NEW Excel column
                "placed":
                    clean(
                        ws.cell(row, 7).value
                    ),

                # NEW Excel positions
                "required":
                    ws.cell(row, 36).value,

                "total_registered":
                    ws.cell(row, 37).value,

                "gender":
                    clean(
                        ws.cell(row, 38).value
                    ),

                # 4-1 is column 39
                "current_4_1":
                    clean(
                        ws.cell(row, 39).value
                    ),

                # Column 40 is the registration flag
                "registered_4_1":
                    ws.cell(row, 40).value,

                # 4-2 is column 41
                "current_4_2":
                    clean(
                        ws.cell(row, 41).value
                    )
            }

            wb.close()

            return student

    wb.close()

    return None


# ============================================================
# SEMESTER CONFIGURATION
# ============================================================

# Column numbers are based on the NEW
# Y23 Skill Certification Details_ 10-08-26.xlsx

SEMESTERS = {

    "2-1": {
        "course_code": 8,
        "course_name": 9,
        "erp_status": 10,
        "current_status": 11,
        "replacement_code": 13,
        "replacement_name": 14
    },

    "2-2": {
        "course_code": 15,
        "course_name": 16,
        "erp_status": 17,
        "current_status": 18,
        "replacement_code": 20,
        "replacement_name": 21
    },

    "3-1": {
        "course_code": 22,
        "course_name": 23,
        "erp_status": 24,
        "current_status": 25,
        "replacement_code": 27,
        "replacement_name": 28
    },

    "3-2": {
        "course_code": 29,
        "course_name": 30,
        "erp_status": 31,
        "current_status": 32,
        "replacement_code": 34,
        "replacement_name": 35
    }
}


# ============================================================
# STUDENT HISTORY
# ============================================================

def get_student_history(student_id):

    student = find_student(student_id)

    if not student:
        return None

    wb = get_workbook()

    ws = wb[REGISTRATION_SHEET]

    row = student["row"]

    history = {}

    for semester, columns in SEMESTERS.items():

        course_code = clean(
            ws.cell(
                row,
                columns["course_code"]
            ).value
        )

        course_name = clean(
            ws.cell(
                row,
                columns["course_name"]
            ).value
        )

        erp_status = clean(
            ws.cell(
                row,
                columns["erp_status"]
            ).value
        )

        current_status = clean(
            ws.cell(
                row,
                columns["current_status"]
            ).value
        )

        replacement_code = clean(
            ws.cell(
                row,
                columns["replacement_code"]
            ).value
        )

        replacement_name = clean(
            ws.cell(
                row,
                columns["replacement_name"]
            ).value
        )

        history[semester] = {

            "course_code":
                course_code,

            "course_name":
                course_name,

            "erp_status":
                erp_status,

            "current_status":
                current_status,

            "replacement_code":
                replacement_code,

            "replacement_name":
                replacement_name,

            "can_replace":
                can_replace(
                    course_code,
                    erp_status,
                    current_status
                )
        }

    wb.close()

    return history


# ============================================================
# REPLACEMENT ELIGIBILITY
# ============================================================

def can_replace(
    course_code,
    erp_status,
    current_status
):

    if not course_code:
        return False

    erp = clean(
        erp_status
    ).lower()

    current = clean(
        current_status
    ).lower()

    # Explicitly completed
    # means replacement is NOT allowed.
    completed_values = {

        "completed",

        "completed and results released",

        "completed but marked as gp",

        "completed but result yet to be released",

        "pass",

        "passed"
    }

    for value in completed_values:

        if current == value:

            return False

    # Explicitly not completed
    # means replacement is allowed.
    if (
        "not completed" in current
        or "not done" in current
    ):

        return True

    # GP in this Excel indicates
    # pending/not completed.
    if erp == "gp":

        return True

    return False


# ============================================================
# ALL PREVIOUSLY REGISTERED COURSES
# ============================================================

def get_registered_codes(student_id):

    student = find_student(student_id)

    if not student:
        return set()

    wb = get_workbook()

    ws = wb[REGISTRATION_SHEET]

    row = student["row"]

    registered = set()

    # --------------------------------------------------------
    # 2-1, 2-2, 3-1, 3-2
    # --------------------------------------------------------

    for columns in SEMESTERS.values():

        # Original certification
        code = clean(
            ws.cell(
                row,
                columns["course_code"]
            ).value
        )

        if code:

            registered.add(code)

        # Replacement certification
        replacement = clean(
            ws.cell(
                row,
                columns["replacement_code"]
            ).value
        )

        if replacement:

            registered.add(replacement)

    # --------------------------------------------------------
    # 4-1
    # --------------------------------------------------------

    current_4_1 = clean(
        ws.cell(row, 39).value
    )

    if current_4_1:

        # Example:
        #
        # 23CC3107-MONGODB ASSOCIATE...
        #
        # Extract code before first "-".

        code = current_4_1.split(
            "-",
            1
        )[0].strip()

        if code:

            registered.add(code)

    # --------------------------------------------------------
    # 4-2
    # --------------------------------------------------------

    current_4_2 = clean(
        ws.cell(row, 41).value
    )

    if current_4_2:

        code = current_4_2.split(
            "-",
            1
        )[0].strip()

        if code:

            registered.add(code)

    wb.close()

    return registered


# ============================================================
# AVAILABLE COURSES
# ============================================================

def get_available_courses(student_id):

    # First verify student exists.
    student = find_student(student_id)

    if not student:
        return []

    registered_codes = (
        get_registered_codes(
            student_id
        )
    )

    courses = get_courses()

    available = []

    for course in courses:

        if course["code"] not in registered_codes:

            available.append(course)

    return available


# ============================================================
# STUDENT DASHBOARD
# ============================================================

def get_student_dashboard(student_id):

    student = find_student(
        student_id
    )

    if not student:

        return None

    history = get_student_history(
        student_id
    )

    available_courses = (
        get_available_courses(
            student_id
        )
    )

    # Semesters where replacement
    # may be required.
    replacement_needed = []

    for semester, data in history.items():

        if data["can_replace"]:

            replacement_needed.append(
                semester
            )

    return {

        "student_id":
            student["student_id"],

        "student_name":
            student["student_name"],

        "placed":
            student["placed"],

        "required":
            student["required"],

        "total_registered":
            student["total_registered"],

        "gender":
            student["gender"],

        "current_4_1":
            student["current_4_1"],

        "registered_4_1":
            student["registered_4_1"],

        "current_4_2":
            student["current_4_2"],

        "history":
            history,

        "replacement_needed":
            replacement_needed,

        "available_courses":
            available_courses
    }


# ============================================================
# REPLACE CERTIFICATION
# ============================================================

def replace_certification(
    student_id,
    semester,
    new_course_code
):

    if semester not in SEMESTERS:

        return {

            "success": False,

            "message":
                "Invalid semester."
        }

    student = find_student(
        student_id
    )

    if not student:

        return {

            "success": False,

            "message":
                "Student ID not found."
        }

    new_course_code = clean(
        new_course_code
    )

    # --------------------------------------------------------
    # Check whether selected course exists
    # --------------------------------------------------------

    courses = get_courses()

    selected_course = None

    for course in courses:

        if course["code"] == new_course_code:

            selected_course = course

            break

    if not selected_course:

        return {

            "success": False,

            "message":
                "Invalid certification selected."
        }

    # --------------------------------------------------------
    # Prevent duplicate certification
    # --------------------------------------------------------

    registered_codes = (
        get_registered_codes(
            student_id
        )
    )

    if new_course_code in registered_codes:

        return {

            "success": False,

            "message":
                "This certification was already "
                "registered by this student."
        }

    wb = get_workbook()

    ws = wb[REGISTRATION_SHEET]

    row = student["row"]

    columns = SEMESTERS[semester]

    current_code = clean(
        ws.cell(
            row,
            columns["course_code"]
        ).value
    )

    current_status = clean(
        ws.cell(
            row,
            columns["current_status"]
        ).value
    )

    erp_status = clean(
        ws.cell(
            row,
            columns["erp_status"]
        ).value
    )

    # --------------------------------------------------------
    # Check replacement eligibility
    # --------------------------------------------------------

    if not can_replace(
        current_code,
        erp_status,
        current_status
    ):

        wb.close()

        return {

            "success": False,

            "message":
                "This certification is not "
                "eligible for replacement."
        }

    # --------------------------------------------------------
    # Prevent duplicate replacement
    # --------------------------------------------------------

    existing_replacement = clean(
        ws.cell(
            row,
            columns["replacement_code"]
        ).value
    )

    if existing_replacement:

        wb.close()

        return {

            "success": False,

            "message":
                "A replacement has already "
                "been registered for this semester."
        }

    # --------------------------------------------------------
    # Write replacement
    # --------------------------------------------------------

    ws.cell(
        row,
        columns["replacement_code"]
    ).value = selected_course["code"]

    ws.cell(
        row,
        columns["replacement_name"]
    ).value = selected_course["title"]

    wb.save(EXCEL_FILE)

    wb.close()

    return {

        "success": True,

        "message":
            f"{semester} certification "
            "replacement submitted successfully.",

        "semester":
            semester,

        "course":
            selected_course
    }


# ============================================================
# 4-1 REGISTRATION
# ============================================================

def register_4_1(
    student_id,
    course_code
):

    student = find_student(
        student_id
    )

    if not student:

        return {

            "success": False,

            "message":
                "Student ID not found."
        }

    course_code = clean(
        course_code
    )

    courses = get_courses()

    selected_course = None

    for course in courses:

        if course["code"] == course_code:

            selected_course = course

            break

    if not selected_course:

        return {

            "success": False,

            "message":
                "Invalid certification selected."
        }

    # --------------------------------------------------------
    # Prevent duplicate certification
    # --------------------------------------------------------

    registered_codes = (
        get_registered_codes(
            student_id
        )
    )

    if course_code in registered_codes:

        return {

            "success": False,

            "message":
                "You have already registered "
                "for this certification."
        }

    wb = get_workbook()

    ws = wb[REGISTRATION_SHEET]

    row = student["row"]

    # 4-1 is column 39
    current_4_1 = clean(
        ws.cell(row, 39).value
    )

    if current_4_1:

        wb.close()

        return {

            "success": False,

            "message":
                "A 4-1 registration already exists."
        }

    value = (
        f"{selected_course['code']}-"
        f"{selected_course['title']}"
    )

    ws.cell(
        row,
        39
    ).value = value

    # Column 40 is the registration flag.
    # Set it to True after successful registration.
    ws.cell(
        row,
        40
    ).value = True

    wb.save(EXCEL_FILE)

    wb.close()

    return {

        "success": True,

        "message":
            "4-1 registration successful.",

        "course":
            selected_course
    }


# ============================================================
# 4-2 REGISTRATION
# ============================================================

def register_4_2(
    student_id,
    course_code
):

    student = find_student(
        student_id
    )

    if not student:

        return {

            "success": False,

            "message":
                "Student ID not found."
        }

    course_code = clean(
        course_code
    )

    courses = get_courses()

    selected_course = None

    for course in courses:

        if course["code"] == course_code:

            selected_course = course

            break

    if not selected_course:

        return {

            "success": False,

            "message":
                "Invalid certification selected."
        }

    # --------------------------------------------------------
    # Prevent duplicate certification
    # --------------------------------------------------------

    registered_codes = (
        get_registered_codes(
            student_id
        )
    )

    if course_code in registered_codes:

        return {

            "success": False,

            "message":
                "You have already registered "
                "for this certification."
        }

    wb = get_workbook()

    ws = wb[REGISTRATION_SHEET]

    row = student["row"]

    # 4-2 is column 41
    current_4_2 = clean(
        ws.cell(row, 41).value
    )

    if current_4_2:

        wb.close()

        return {

            "success": False,

            "message":
                "A 4-2 registration already exists."
        }

    value = (
        f"{selected_course['code']}-"
        f"{selected_course['title']}"
    )

    ws.cell(
        row,
        41
    ).value = value

    wb.save(EXCEL_FILE)

    wb.close()

    return {

        "success": True,

        "message":
            "4-2 registration successful.",

        "course":
            selected_course
    }

