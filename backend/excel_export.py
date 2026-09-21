from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db import courses_collection, students_collection
from excel_import import COURSE_SHEET, REGISTRATION_SHEET, SEMESTERS

# Styling mirrored from the master workbook: light-grey header
# (theme 0, tint -0.15), bold black text, thin borders on every cell.
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FONT = Font(bold=True, size=11)
_side = Side(style="thin")
THIN = Border(left=_side, right=_side, top=_side, bottom=_side)

NO_FILL_COLUMNS = {5, 40}
WRAP_COLUMNS = {10, 11, 12, 17, 18, 19, 24, 25, 26, 31, 32, 33, 36}

COLUMN_WIDTHS = [
    8.8, 17.8, 16.5, 36.7, 13.3, 36.7, 7.7, 11.7, 47.2, 13.2,
    14.5, 43.7, 27.2, 31.9, 11.3, 39.3, 7.5, 18.0, 29.3, 27.2,
    31.9, 11.0, 35.2, 7.3, 23.0, 27.2, 27.2, 31.9, 11.7, 45.7,
    7.8, 14.8, 27.2, 27.2, 31.9, 18.5, 25.5, 11.8, 56.7, 12.0, 23.2
]

ERP_HEADER = (
    "ERP Status\n"
    "1. Completed and results released\n"
    "2. Completed but marked as GP\n"
    "3. Not done/GP\n"
    "4. Completed but result yet to be released"
)

# Column 1..41, matching the master workbook exactly.
HEADERS = [
    "S.No",
    "Faculty Incharge",
    "STUDENT_ID",
    "STUDENT NAME",
    "COUNSELLOR EMP ID",
    "Counsellor Name",
    "PLACED",
]

for _sem in SEMESTERS:
    HEADERS += [
        _sem,
        f"{_sem} Course title",
        ERP_HEADER,
        "Student's Current Status",
        "Verifiable link",
        f"{_sem} REPLACEMENT COURSE CODE",
        f"{_sem} REPLACEMENT COURSE NAME",
    ]

HEADERS += [
    "Required",
    "Total Courses registered",
    "Gender",
    "4-1 Course to be registered",
    None,
    "4-2 Course to be registered",
]


def is_placed(value):
    return (value or "").strip().lower() == "placed"


def build_report(scope="all"):
    """
    Rebuild the master workbook from MongoDB, in the exact upload format
    so the export can be edited and re-uploaded unchanged.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = REGISTRATION_SHEET

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = HEADER_FONT
        cell.border = THIN
        cell.alignment = Alignment(
            horizontal="left" if col == 38 else "center",
            vertical="center",
            wrap_text=col in WRAP_COLUMNS
        )
        if col not in NO_FILL_COLUMNS:
            cell.fill = HEADER_FILL

    ws.row_dimensions[1].height = 31.5
    ws.freeze_panes = "E2"

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    query = {} if scope == "all" else {
        "placed": {"$regex": r"^\s*placed\s*$", "$options": "i"}
    }

    row = 1

    for student in students_collection.find(query).sort("student_id", 1):

        row += 1
        history = student.get("history", {})

        ws.cell(row, 1).value = row - 1
        ws.cell(row, 2).value = student.get("faculty_incharge", "")
        ws.cell(row, 3).value = student.get("student_id", "")
        ws.cell(row, 4).value = student.get("student_name", "")
        ws.cell(row, 5).value = student.get("counsellor_emp_id", "")
        ws.cell(row, 6).value = student.get("counsellor_name", "")
        ws.cell(row, 7).value = student.get("placed", "")

        for semester, columns in SEMESTERS.items():
            data = history.get(semester, {})
            for field, col in columns.items():
                ws.cell(row, col).value = data.get(field, "")

        ws.cell(row, 36).value = student.get("required")
        ws.cell(row, 37).value = student.get("total_registered")
        ws.cell(row, 38).value = student.get("gender", "")
        ws.cell(row, 39).value = student.get("current_4_1", "")
        ws.cell(row, 40).value = bool(student.get("registered_4_1"))
        ws.cell(row, 41).value = student.get("current_4_2", "")

        for col in range(1, 42):
            ws.cell(row, col).border = THIN

    # Included so an exported report is valid input for re-upload.
    cs = wb.create_sheet(COURSE_SHEET)
    cs.append(["Course Code", "Course Title"])
    for course in courses_collection.find({}, {"_id": 0}).sort("code", 1):
        cs.append([course["code"], course["title"]])

    for cell in cs[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN

    cs.column_dimensions["A"].width = 16.5
    cs.column_dimensions["B"].width = 56.7

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer, row - 1
